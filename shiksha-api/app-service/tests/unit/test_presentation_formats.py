import io
import shutil
import zipfile

import aiofiles
import pytest
from fastapi import UploadFile
from pptx import Presentation

from app.routers.presentation import ALLOWED_MIMES
from app.services.presentation import docparser
from app.services.presentation.docparser import OFFICE_MIME_SUFFIXES
from app.services.presentation.utils import LibreOffice, save_file_with_hash
from app.utils.storage import Storage

PDF = "application/pdf"
DOC = "application/msword"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPT = "application/vnd.ms-powerpoint"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
OFFICE = [DOC, DOCX, PPT, PPTX]

needs_soffice = pytest.mark.skipif(shutil.which("soffice") is None, reason="soffice not installed")


def make_pptx(text: str) -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = text
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def make_docx() -> bytes:
    # minimal OOXML package, enough for libmagic to identify it as docx
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
        z.writestr("_rels/.rels", '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
        z.writestr("word/document.xml", '<?xml version="1.0"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Photosynthesis</w:t></w:r></w:p></w:body></w:document>')
    return buf.getvalue()


async def reader(tmp_path, data: bytes):
    path = tmp_path / "in.bin"
    path.write_bytes(data)
    return await aiofiles.open(path, "rb")


class TestAllowedMimes:
    @pytest.mark.parametrize("mime", [PDF, "text/plain", "text/markdown", *OFFICE])
    def test_accepts(self, mime):
        assert mime in ALLOWED_MIMES

    def test_every_non_text_mime_has_a_conversion_path(self):
        # An allowed mime that cannot be extracted to markdown would be sent raw to OpenAI and fail
        # with "Unsupported binary content type" - this is the drift this test guards against.
        convertible = {PDF, *OFFICE_MIME_SUFFIXES}
        assert {m for m in ALLOWED_MIMES if not m.startswith("text/")} <= convertible

    def test_office_suffixes_match_mimes(self):
        assert OFFICE_MIME_SUFFIXES == {DOC: ".doc", DOCX: ".docx", PPT: ".ppt", PPTX: ".pptx"}


class TestLibreOfficeInputSuffix:
    @pytest.mark.parametrize("suffix", [".doc", ".docx", ".ppt", ".pptx"])
    async def test_input_file_keeps_original_suffix(self, monkeypatch, suffix):
        lo = LibreOffice()
        lo.path = "soffice"
        seen = {}

        async def fake_call(*args):
            in_path = args[-1]
            seen["in"] = in_path
            outdir = args[args.index("--outdir") + 1]
            with open(f"{outdir}/input.pdf", "wb") as f:
                f.write(b"%PDF-fake")
            return b""

        monkeypatch.setattr(lo, "_call", fake_call)
        out = await lo.convert(io.BytesIO(b"data"), "pdf", suffix)
        assert seen["in"].endswith("input" + suffix)
        assert out == b"%PDF-fake"

    async def test_default_suffix_is_pptx(self, monkeypatch):
        lo = LibreOffice()
        seen = {}

        async def fake_call(*args):
            seen["in"] = args[-1]
            with open(f"{args[args.index('--outdir') + 1]}/input.pptx", "wb") as f:
                f.write(b"x")
            return b""

        monkeypatch.setattr(lo, "_call", fake_call)
        await lo.convert(io.BytesIO(b"data"), "pptx")
        assert seen["in"].endswith("input.pptx")


class TestTransform:
    @pytest.mark.parametrize("mime", OFFICE)
    async def test_office_converted_to_pdf_then_markdown(self, tmp_path, monkeypatch, mime):
        calls = {}

        async def fake_convert(self, data, output_format, input_suffix=".pptx"):
            calls["args"] = (data.read(), output_format, input_suffix)
            return b"%PDF-converted"

        monkeypatch.setattr(LibreOffice, "convert", fake_convert)
        monkeypatch.setattr(docparser, "_to_markdown", lambda data: "# md from " + data.decode())

        f = await reader(tmp_path, b"office-bytes")
        try:
            result = await docparser._transform(f, mime)
        finally:
            await f.close()

        assert result == "# md from %PDF-converted"
        assert calls["args"] == (b"office-bytes", "pdf", OFFICE_MIME_SUFFIXES[mime])

    async def test_pdf_goes_straight_to_markdown(self, tmp_path, monkeypatch):
        async def boom(*a, **k):
            raise AssertionError("LibreOffice must not run for PDF")

        monkeypatch.setattr(LibreOffice, "convert", boom)
        monkeypatch.setattr(docparser, "_to_markdown", lambda data: "pdf:" + data.decode())
        f = await reader(tmp_path, b"pdfbytes")
        try:
            assert await docparser._transform(f, PDF) == "pdf:pdfbytes"
        finally:
            await f.close()

    @pytest.mark.parametrize("mime", ["text/plain", "text/markdown"])
    async def test_text_is_passed_through(self, tmp_path, mime):
        f = await reader(tmp_path, b"hello")
        try:
            assert await docparser._transform(f, mime) is None
        finally:
            await f.close()

    async def test_conversion_failure_propagates(self, tmp_path, monkeypatch):
        async def fail(self, *a, **k):
            raise RuntimeError("soffice crashed")

        monkeypatch.setattr(LibreOffice, "convert", fail)
        f = await reader(tmp_path, b"x")
        try:
            with pytest.raises(RuntimeError, match="soffice crashed"):
                await docparser._transform(f, DOCX)
        finally:
            await f.close()

    async def test_transform_writes_markdown_next_to_out_path(self, tmp_path, monkeypatch):
        async def fake_convert(self, data, output_format, input_suffix=".pptx"):
            return b"%PDF"

        monkeypatch.setattr(LibreOffice, "convert", fake_convert)
        monkeypatch.setattr(docparser, "_to_markdown", lambda data: "# Converted")
        storage = Storage("memory", "")
        f = await reader(tmp_path, b"x")
        try:
            name = await docparser.transform(storage, f, DOCX, "out/stem/stem")
        finally:
            await f.close()
        assert name == "stem.md"
        assert await storage.read_text("out/stem/stem.md") == "# Converted"

    @needs_soffice
    @pytest.mark.parametrize("builder,mime", [(lambda: make_pptx("Photosynthesis basics"), PPTX)])
    async def test_real_libreoffice_conversion(self, tmp_path, builder, mime):
        f = await reader(tmp_path, builder())
        try:
            md = await docparser._transform(f, mime)
        finally:
            await f.close()
        assert md is not None and "Photosynthesis" in md


class TestUploadMimeDetection:
    @pytest.fixture
    def storage(self):
        return Storage("memory", "")

    @staticmethod
    def upload(data: bytes, name: str) -> UploadFile:
        return UploadFile(io.BytesIO(data), filename=name)

    async def test_pptx_accepted(self, storage):
        name, mime = await save_file_with_hash(storage, self.upload(make_pptx("t"), "deck.pptx"), "deck.pptx", ALLOWED_MIMES, 1024, 10_000_000)
        assert mime == PPTX and name.endswith(".pptx")

    async def test_docx_accepted(self, storage):
        name, mime = await save_file_with_hash(storage, self.upload(make_docx(), "book.docx"), "book.docx", ALLOWED_MIMES, 1024, 10_000_000)
        assert mime == DOCX and name.endswith(".docx")

    async def test_unsupported_binary_rejected(self, storage):
        png = b"\x89PNG\r\n\x1a\n" + b"\x00" * 32
        with pytest.raises(ValueError, match="Unexpected mime"):
            await save_file_with_hash(storage, self.upload(png, "x.png"), "x.png", ALLOWED_MIMES, 1024, 10_000_000)

    async def test_oversize_rejected(self, storage):
        with pytest.raises(ValueError, match="Bad file size"):
            await save_file_with_hash(storage, self.upload(b"a" * 100, "a.txt"), "a.txt", ALLOWED_MIMES, 10, 50)
