import asyncio
import io
import json
import logging
import re
import time
from collections.abc import Awaitable, Callable
from contextlib import nullcontext
from dataclasses import dataclass
from datetime import datetime, timezone
from enum import StrEnum
from functools import cache
from graphlib import CycleError, TopologicalSorter
from textwrap import dedent
from typing import Any, Literal

import openai
from langfuse import get_client as _get_client
from omni_ingest.core.model import Step, StepResult, StepStatus
from omni_ingest.core.pipeline import IngestionContext, register_step
from openpyxl import load_workbook
from pydantic import (
    AliasChoices,
    BaseModel,
    Field,
    InstanceOf,
    PrivateAttr,
    ValidationError,
    field_validator,
    model_validator,
)
from qdrant_client import AsyncQdrantClient
from qdrant_client.http.models import FieldCondition, Filter, MatchValue

import period_plan_legacy_rag as legacy

log = logging.getLogger(__name__)
get_client = cache(_get_client)
PREVIOUS = "Previously Generated Section: "
FENCE = re.compile(r"^\s*```[A-Za-z]*\s*|\s*```\s*$")
PLACEHOLDER = re.compile(r"\$\{([A-Za-z0-9_]+)\}")
RETRY_ATTEMPTS = 3
MIN_DESCRIPTION_CHARS = 10
MIN_SUMMARY_CHARS = 100
SUMMARY_LIMIT = 800
QUERY_LIMIT = 500
CUT_RATIO = 0.8


@dataclass
class Connections:
    qdrant: AsyncQdrantClient
    llm: openai.AsyncAzureOpenAI
    post_status: Callable[[dict], Awaitable[None]]
    is_done: Callable[[str], bool]
    save: Callable[[str, dict], None]


class LPLevel(StrEnum):
    CHAPTER = "CHAPTER"
    SUBTOPIC = "SUBTOPIC"
    TELANGANA_ENGLISH_RESOURCE_PLAN = "TELANGANA_ENGLISH_RESOURCE_PLAN"


class PlanStatus(StrEnum):
    SKIPPED = "skipped"
    FAILED = "failed"
    COMPLETED = "completed"


class EnglishLPType(StrEnum):
    PROSE = "PROSE"
    POEM = "POEM"
    NONE = "NONE"


class PlanResult(BaseModel):
    status: PlanStatus
    plan: dict | None = None
    error: str | None = None


class Dependency(BaseModel):
    section_id: str


class SectionDef(BaseModel):
    id: str
    title: str
    description: str
    mode: Literal["rag", "gpt"]
    dependencies: list[Dependency] = []
    output_format: dict[str, Any] | None = None


class Workflow(BaseModel):
    id: str = Field(alias="_id")
    sections: list[SectionDef] = Field(min_length=1)

    @model_validator(mode="after")
    def _check_dependencies(self) -> "Workflow":
        ids = {s.id for s in self.sections}
        for s in self.sections:
            for d in s.dependencies:
                if d.section_id not in ids:
                    raise ValueError(
                        f"Section '{s.id}' depends on '{d.section_id}', which "
                        "is not a section in the workflow."
                    )
        graph = {
            s.id: [d.section_id for d in s.dependencies] for s in self.sections
        }
        try:
            TopologicalSorter(graph).prepare()
        except CycleError as e:
            raise ValueError(
                "Detected a cycle in the workflow dependencies: "
                + " -> ".join(e.args[1])
            ) from e
        return self


class QdrantIndex(BaseModel):
    collection: str
    key: str
    value: str


def parse_index(path: str) -> QdrantIndex | legacy.BlobIndex:
    graph = "graph-index" in path
    if path.startswith("qdrant/") and not graph:
        try:
            _, collection, flt = path.split("/", 2)
            key, value = flt.split(":", 1)
        except ValueError:
            raise ValueError(
                f"index_path '{path}' must look like "
                "qdrant/<collection>/<key>:<value>"
            ) from None
        return QdrantIndex(collection=collection, key=key, value=value)
    container, sep, prefix = path.partition("/")
    if not sep:
        raise ValueError(
            f"Chapter index '{path}' must look like <container>/<path>."
        )
    return legacy.BlobIndex(
        path, "graph" if graph else "vector", container, prefix
    )


class ChapterInfo(BaseModel):
    id: str
    index_path: str
    chapter_title: str = Field(min_length=1)
    _index: QdrantIndex | legacy.BlobIndex = PrivateAttr()

    @model_validator(mode="after")
    def _parse_index(self) -> "ChapterInfo":
        self._index = parse_index(self.index_path)
        return self

    @property
    def index(self) -> QdrantIndex | legacy.BlobIndex:
        return self._index


class Section(BaseModel):
    section_id: str
    section_title: str
    content: Any
    regen_feedback: str | None = None


class LessonPlanContent(BaseModel):
    sections: list[Section] = []


class LessonPlanGenerationInput(BaseModel):
    user_id: str = "ADMIN"
    workflow: Workflow
    lp_id: str = Field(validation_alias=AliasChoices("lp_id", "_id"))
    chapter_info: ChapterInfo
    lp_level: LPLevel
    learning_outcomes: list[str]
    lp_type_english: EnglishLPType = EnglishLPType.NONE
    start_from_section_id: str | None = None
    subtopics: list[str] = []
    lesson_plan: LessonPlanContent | None = None
    prompt_variables: dict[str, str] = {}
    additional_context: str | None = None
    created_at: int = Field(default_factory=lambda: int(time.time()))

    @field_validator("additional_context")
    @classmethod
    def _strip_context(cls, value: str | None) -> str | None:
        return (value or "").strip() or None

    @model_validator(mode="after")
    def _need_subtopics(self) -> "LessonPlanGenerationInput":
        if self.lp_level == LPLevel.SUBTOPIC and not self.subtopics:
            raise ValueError(
                "Subtopics must be provided for SUBTOPIC level lesson plans"
            )
        return self

    @model_validator(mode="after")
    def _check_resume(self) -> "LessonPlanGenerationInput":
        sid = self.start_from_section_id
        if not sid:
            return self
        sections = {s.id: s for s in self.workflow.sections}
        if sid not in sections:
            raise ValueError(f"Section ID '{sid}' not found in workflow")
        if self.lesson_plan is None:
            return self
        content = {s.section_id: s.content for s in self.lesson_plan.sections}
        if empty := [k for k, v in content.items() if v is None]:
            raise ValueError(
                f"The lesson plan has no content for sections {empty}. Send "
                "the saved content of each section, or leave the section out."
            )
        required = [d.section_id for d in sections[sid].dependencies]
        if missing := [d for d in required if d not in content]:
            raise ValueError(
                f"Missing dependency content for section '{sid}'. Required "
                f"dependencies not provided: {missing}"
            )
        return self

    @property
    def variant(self) -> Literal["fresh", "regen", "resource"]:
        if self.lesson_plan is not None and not self.start_from_section_id:
            return "regen"
        if self.lp_level == LPLevel.TELANGANA_ENGLISH_RESOURCE_PLAN:
            return "resource"
        return "fresh"


@dataclass
class Node:
    title: str
    deps: list[str]
    out: Any = None
    done: bool = False


def reason(e: Exception) -> str:
    if isinstance(e, ValidationError):
        return "; ".join(
            f"{'.'.join(map(str, x['loc']))}: {x['msg']}" for x in e.errors()
        )
    return str(e) or type(e).__name__


def parse_json(text: str, sec: SectionDef | None = None) -> Any:
    try:
        return json.loads(FENCE.sub("", text.strip()))
    except ValueError as e:
        if sec and sec.output_format:
            raise ValueError(
                f"The model reply for section '{sec.id}' is not valid JSON "
                f"({e}), but the section has an output_format. It starts with: "
                f"{text[:80]}"
            ) from e
        log.warning(
            "The model reply for section '%s' is not JSON. Keeping it as "
            "text. It starts with: %.80s",
            sec.id if sec else "?",
            text,
        )
        return text


def parse_reply(text: str, sec: SectionDef) -> Any:
    if not text:
        raise ValueError(
            f"The model returned an empty answer for section '{sec.id}'."
        )
    return parse_json(text, sec)


def dump(content: Any) -> str:
    return json.dumps(content, ensure_ascii=False)


def build_nodes(inp: LessonPlanGenerationInput) -> dict[str, Node]:
    nodes = {
        s.id: Node(s.title, [d.section_id for d in s.dependencies])
        for s in inp.workflow.sections
    }
    if inp.lesson_plan is None:
        return nodes
    if inp.variant == "regen":
        for s in inp.lesson_plan.sections:
            if s.section_id in nodes:
                other = f"other_{s.section_id}"
                note = (
                    f"\n\n**Feedback for regeneration: {s.regen_feedback}**"
                    if s.regen_feedback
                    else ""
                )
                nodes[other] = Node(
                    PREVIOUS + s.section_title,
                    [],
                    out=dump(s.content) + note,
                    done=True,
                )
                nodes[s.section_id].deps.append(other)
    elif inp.start_from_section_id:
        for s in inp.lesson_plan.sections:
            if s.section_id in nodes:
                nodes[s.section_id].out = s.content
                nodes[s.section_id].done = True
    return nodes


async def run_graph(
    nodes: dict[str, Node],
    generate: Callable[[str, dict[str, Any]], Awaitable[Any]],
) -> None:
    TopologicalSorter({k: n.deps for k, n in nodes.items()}).prepare()
    tasks: dict[str, asyncio.Task] = {}

    async def run(k: str) -> None:
        for d in nodes[k].deps:
            if d in tasks:
                await tasks[d]
        deps = {nodes[d].title: nodes[d].out for d in nodes[k].deps}
        nodes[k].out = await generate(k, deps)

    try:
        async with asyncio.TaskGroup() as group:
            for k, n in nodes.items():
                if not n.done:
                    tasks[k] = group.create_task(run(k))
    except BaseExceptionGroup as e:
        for other in e.exceptions[1:]:
            log.error("Another section also failed.", exc_info=other)
        raise e.exceptions[0] from None


def fill(template: str, **slots: str) -> str:
    return re.sub(r"\n{3,}", "\n\n", template.format(**slots)).strip()


def truncate(text: str, limit: int) -> str:
    if len(text) <= limit:
        return text
    cut = text[:limit]
    keep = limit * CUT_RATIO
    if (period := cut.rfind(".")) > keep:
        return cut[: period + 1]
    if (space := cut.rfind(" ")) > keep:
        return cut[:space]
    return cut


def default_query(inp: LessonPlanGenerationInput, regen: bool) -> str:
    sep = "\n\n" if regen else "\n"
    outcomes = "\n".join(inp.learning_outcomes)
    if inp.lp_level == LPLevel.SUBTOPIC:
        label = "Topic Titles" if regen else "Topic(s)"
        return (
            f"{label}: {'; '.join(inp.subtopics)}{sep}"
            f"Learning Outcomes: {outcomes}"
        )
    return (
        f"Chapter Title: {inp.chapter_info.chapter_title}{sep}"
        f"Learning Outcomes: {outcomes}"
    )


def clean_query(query: str, limit: int = QUERY_LIMIT) -> str:
    query = query.strip()
    if not query:
        raise ValueError("Generated query is empty")
    for prefix in (
        "Retrieval Query:",
        "Query:",
        "Search Query:",
        "Generated Query:",
    ):
        if query.startswith(prefix):
            query = query[len(prefix) :].strip()
    return truncate(query, limit)


class QdrantRetrieve(BaseModel):
    embed_model: str
    top_k: int = 10
    vector_name: str = "text-dense"

    async def search(
        self, conn: Connections, index: QdrantIndex, query: str
    ) -> list[str]:
        flt = Filter(
            must=[
                FieldCondition(
                    key=index.key, match=MatchValue(value=index.value)
                )
            ]
        )
        vector = (
            (
                await conn.llm.embeddings.create(
                    model=self.embed_model, input=query
                )
            )
            .data[0]
            .embedding
        )
        result = await conn.qdrant.query_points(
            collection_name=index.collection,
            query=vector,
            using=self.vector_name,
            query_filter=flt,
            limit=self.top_k,
            with_payload=True,
        )
        if not result.points:
            raise ValueError(
                f"Qdrant returned no passages for collection "
                f"'{index.collection}' with {index.key}:{index.value}. Check "
                "that the chapter was indexed."
            )
        texts = []
        for p in result.points:
            text = (
                json.loads(p.payload["_node_content"])["text"]
                if "_node_content" in p.payload
                else p.payload.get("text")
            )
            if not text:
                raise ValueError(
                    f"Qdrant point {p.id} has no text payload. Check how the "
                    "chapter was indexed."
                )
            texts.append(text)
        return texts


class VariantPrompts(BaseModel):
    rag: str
    json_format: str = Field(alias="json")
    md: str
    template: str


class FreshPrompts(VariantPrompts):
    subtopic: str
    chapter: str
    outcomes: str
    deps: str


class RegenPrompts(VariantPrompts):
    subtopic: str
    chapter: str
    deps_current: str
    deps_previous: str


class ResourcePrompts(VariantPrompts):
    intro: str
    deps: str


class Prompts(BaseModel):
    system: str
    summarize: str
    retrieval_query: str
    rag_answer: str
    journey: str
    additional: str
    fresh: FreshPrompts
    regen: RegenPrompts
    resource: ResourcePrompts


class SectionGraph(BaseModel, Step):
    model: str
    prompts: Prompts
    retrieve: QdrantRetrieve
    conn: InstanceOf[Connections] | None = Field(default=None, exclude=True)
    webhook: bool = True
    force_mode: Literal["", "rag", "gpt"] = ""
    concurrency: int = 1

    async def _chat(
        self, content: str, system: bool = True, **kwargs: Any
    ) -> str:
        messages = (
            [{"role": "system", "content": self.prompts.system}]
            if system
            else []
        ) + [{"role": "user", "content": content}]
        response = await self.conn.llm.chat.completions.create(
            model=self.model, messages=messages, **kwargs
        )
        return (response.choices[0].message.content or "").strip()

    async def _summarize(self, text: str, limit: int = SUMMARY_LIMIT) -> str:
        if len(text.strip()) < MIN_SUMMARY_CHARS:
            return text
        try:
            summary = await self._chat(
                self.prompts.summarize.format(max=limit, text=text),
                temperature=0,
            )
        except openai.APIError as e:
            log.error(
                "Context summary failed (%s). Using the truncated original.",
                e,
                exc_info=True,
            )
            return truncate(text, limit)
        return truncate(summary, limit).strip()

    async def _retrieval_query(
        self, inp: LessonPlanGenerationInput, sec: SectionDef, regen: bool
    ) -> str:
        if len(sec.description.strip()) < MIN_DESCRIPTION_CHARS:
            return default_query(inp, regen)
        topics = (
            inp.subtopics
            if inp.lp_level == LPLevel.SUBTOPIC
            else [inp.chapter_info.chapter_title]
        )
        outcomes = inp.learning_outcomes
        extra = "\nTopics/Subtopics:\n" + "\n".join(f"- {t}" for t in topics)
        if outcomes:
            extra += "\nLearning Outcomes:\n" + "\n".join(
                f"- {o}" for o in outcomes
            )
        try:
            return clean_query(
                await self._chat(
                    self.prompts.retrieval_query.format(
                        description=sec.description,
                        extra=extra,
                        max=QUERY_LIMIT,
                    ),
                    temperature=0,
                )
            )
        except (openai.APIError, ValueError) as e:
            log.warning(
                "Retrieval query generation failed (%s). Using the basic "
                "query.",
                e,
                exc_info=True,
            )
            parts = [sec.description, f"topics: {' '.join(topics)}"]
            if outcomes:
                parts.append(f"learning outcomes: {' '.join(outcomes)}")
            return " ".join(parts)[:QUERY_LIMIT]

    @staticmethod
    def _variable(inp: LessonPlanGenerationInput, name: str) -> str:
        if name not in inp.prompt_variables:
            raise ValueError(
                f"The prompt uses the placeholder ${{{name}}}, but "
                f"prompt_variables has no value for it. Add '{name}' to "
                "prompt_variables."
            )
        return inp.prompt_variables[name]

    def _synthesis(
        self,
        inp: LessonPlanGenerationInput,
        sec: SectionDef,
        deps: dict[str, Any],
        mode: Literal["rag", "gpt"],
    ) -> str:
        variant = inp.variant
        variant_prompts = getattr(self.prompts, variant)
        body = "\n\n".join(
            f"# Section: {t}\n{dump(c)}" for t, c in deps.items()
        )

        def subst(text: str) -> str:
            return PLACEHOLDER.sub(
                lambda m: self._variable(inp, m.group(1)), text
            )

        slots = {
            "journey": self.prompts.journey.format(
                sections=", ".join(s.title for s in inp.workflow.sections),
                current=sec.title,
            ),
            "additional": self.prompts.additional.format(
                text=inp.additional_context
            )
            if inp.additional_context
            else "",
            "description": subst(dedent(sec.description)),
            "rag": variant_prompts.rag if mode == "rag" else "",
            "format": variant_prompts.json_format.format(
                schema=subst(
                    json.dumps(sec.output_format, ensure_ascii=False, indent=2)
                )
            )
            if sec.output_format
            else variant_prompts.md,
            "deps": variant_prompts.deps.format(deps=body)
            if deps and variant != "regen"
            else "",
        }
        topic = "subtopic" if inp.lp_level == LPLevel.SUBTOPIC else "chapter"
        if variant == "fresh":
            query = default_query(inp, False).strip()
            slots["intro"] = fill(
                getattr(variant_prompts, topic), title=sec.title
            )
            slots["outcomes"] = (
                variant_prompts.outcomes.format(query=query) if query else ""
            )
        elif variant == "regen":
            slots["intro"] = getattr(variant_prompts, topic).format(
                title=sec.title, query=default_query(inp, True)
            )
            current = "".join(
                f"# Section: {t}\n{dump(c)}\n\n"
                for t, c in deps.items()
                if not t.startswith(PREVIOUS)
            )
            previous = [
                f"# {t}\n{c}" for t, c in deps.items() if t.startswith(PREVIOUS)
            ]
            if current:
                slots["deps"] = variant_prompts.deps_current.format(
                    deps=current
                )
            if previous:
                slots["deps"] += "\n\n" + variant_prompts.deps_previous.format(
                    previous=previous[-1]
                )
        else:
            slots["intro"] = variant_prompts.intro.format(title=sec.title)
        return fill(variant_prompts.template, **slots)

    async def _rag(
        self,
        inp: LessonPlanGenerationInput,
        sec: SectionDef,
        prompt: str,
        index: QdrantIndex,
    ) -> Any:
        query = await self._retrieval_query(inp, sec, inp.variant == "regen")
        options = (
            {"response_format": {"type": "json_object"}}
            if sec.output_format
            else {}
        )
        context = "\n\n".join(
            await self.retrieve.search(self.conn, index, query)
        )
        answer_prompt = self.prompts.rag_answer.format(
            context=context, query=prompt
        )
        for attempt in range(RETRY_ATTEMPTS):
            try:
                text = await self._chat(
                    answer_prompt, system=False, temperature=0.1, **options
                )
                return parse_reply(text, sec)
            except (openai.APIError, ValueError) as e:
                log.warning(
                    "Attempt %d of %d failed for section '%s' (%s): %s",
                    attempt + 1,
                    RETRY_ATTEMPTS,
                    sec.id,
                    type(e).__name__,
                    e,
                    exc_info=True,
                )
                if attempt == RETRY_ATTEMPTS - 1:
                    raise
                await asyncio.sleep(2**attempt)

    async def _generate(
        self,
        inp: LessonPlanGenerationInput,
        sec: SectionDef,
        deps: dict[str, Any],
    ) -> Any:
        mode = self.force_mode or sec.mode
        prompt = self._synthesis(inp, sec, deps, mode)
        index = inp.chapter_info.index
        if mode == "gpt":
            text = await self._chat(prompt, temperature=0)
        elif isinstance(index, legacy.BlobIndex):
            text = await legacy.answer(
                index, prompt, self.model, self.retrieve.embed_model
            )
        else:
            return await self._rag(inp, sec, prompt, index)
        return parse_reply(text, sec)

    async def _section(
        self,
        inp: LessonPlanGenerationInput,
        sec: SectionDef,
        deps: dict[str, Any],
    ) -> Any:
        with get_client().start_as_current_span(
            name=f"section:{sec.id}",
            input={"mode": self.force_mode or sec.mode},
        ) as span:
            content = await self._generate(inp, sec, deps)
            if content is None:
                raise ValueError(
                    "The model returned no content (JSON null) for section "
                    f"'{sec.id}'."
                )
            span.update(output=content)
            return content

    async def plan(self, inp: LessonPlanGenerationInput) -> PlanResult:
        if self.conn is None:
            raise RuntimeError(
                "SectionGraph has no connections. Build it with "
                "conn=Connections(...)."
            )
        index = inp.chapter_info.index
        langfuse = get_client()
        try:
            async with (
                legacy.session(index)
                if isinstance(index, legacy.BlobIndex)
                else nullcontext()
            ):
                with langfuse.start_as_current_span(
                    name="Shiksha-LP",
                    input=inp.model_dump(mode="json", by_alias=True),
                ) as span:
                    span.update_trace(
                        user_id=inp.user_id, tags=["flow:lesson-plan"]
                    )
                    result = await self._run_plan(inp)
                    span.update(
                        output=result.model_dump(
                            mode="json", include={"status", "error"}
                        )
                    )
            return result
        finally:
            await asyncio.to_thread(langfuse.flush)

    async def _run_plan(self, inp: LessonPlanGenerationInput) -> PlanResult:
        plan_id = inp.lp_id
        request = inp.model_dump(mode="json", by_alias=True)

        async def send(status: str, output: Any = None) -> None:
            if not self.webhook:
                return
            try:
                await self.conn.post_status(
                    {
                        "instance_id": plan_id,
                        "status": status,
                        "timestamp": datetime.now(timezone.utc).isoformat(),
                        "input": request,
                        "output": output,
                    }
                )
            except Exception as e:
                log.error(
                    "Webhook post of status %s failed for plan %s (%s). "
                    "Check WEBHOOK_URL and that the receiver is reachable.",
                    status,
                    plan_id,
                    e,
                    exc_info=True,
                )

        if self.conn.is_done(plan_id):
            return PlanResult(status=PlanStatus.SKIPPED)
        await send("PENDING")
        await send("RUNNING")
        try:
            if inp.additional_context:
                inp.additional_context = await self._summarize(
                    inp.additional_context
                )
            nodes = build_nodes(inp)
            sections = {s.id: s for s in inp.workflow.sections}
            await run_graph(
                nodes, lambda k, deps: self._section(inp, sections[k], deps)
            )
            plan = {
                "_id": plan_id,
                "created_at": inp.created_at,
                "workflow_id": inp.workflow.id,
                "chapter_id": inp.chapter_info.id,
                "subtopics": inp.subtopics,
                "learning_outcomes": inp.learning_outcomes,
                "lp_level": inp.lp_level.value,
                "lp_type_english": inp.lp_type_english.value,
                "sections": [
                    {
                        "section_id": s.id,
                        "section_title": s.title,
                        "content": nodes[s.id].out,
                    }
                    for s in inp.workflow.sections
                ],
            }
            self.conn.save(plan_id, plan)
        except asyncio.CancelledError:
            await send("FAILED", "The run was terminated.")
            raise
        except Exception as e:
            log.exception("Plan %s failed", plan_id)
            error = reason(e)
            await send("FAILED", error)
            return PlanResult(status=PlanStatus.FAILED, error=error)
        await send("COMPLETED", plan)
        return PlanResult(status=PlanStatus.COMPLETED, plan=plan)

    async def run(self, ctx: IngestionContext) -> StepResult:
        payloads = ctx.metadata.get("payloads")
        if not payloads:
            return StepResult(
                status=StepStatus.FAILURE,
                error=(
                    "No payloads found in metadata. Check that the transform "
                    "step before section_graph produced payloads."
                ),
            )
        inputs: list[LessonPlanGenerationInput] = []
        failed: dict[str, PlanResult] = {}
        for i, raw in enumerate(payloads):
            try:
                inputs.append(LessonPlanGenerationInput.model_validate(raw))
            except ValidationError as e:
                ids = raw if isinstance(raw, dict) else {}
                key = ids.get("lp_id") or ids.get("_id") or str(i)
                failed[key] = PlanResult(
                    status=PlanStatus.FAILED, error=reason(e)
                )
        gate = asyncio.Semaphore(self.concurrency)

        async def one(inp: LessonPlanGenerationInput) -> PlanResult:
            async with gate:
                return await self.plan(inp)

        results = await asyncio.gather(*(one(inp) for inp in inputs))
        ctx.metadata["lesson_plans"] = failed | {
            inp.lp_id: r for inp, r in zip(inputs, results, strict=True)
        }
        return StepResult(
            status=StepStatus.SUCCESS,
            metadata={"plans": len(inputs) + len(failed)},
        )


XLSX_HEADERS = (
    "chapter_number",
    "chapter_title",
    "period",
    "topic",
    "learning_outcomes",
)


def cell_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def whole_number(value: Any) -> Any:
    return (
        int(value) if isinstance(value, float) and value.is_integer() else value
    )


def order(value: Any) -> tuple:
    if value is None:
        return (2, 0, "")
    if isinstance(value, (int, float)):
        return (0, value, "")
    return (1, 0, str(value))


class XlsxChapters(BaseModel, Step):
    grade: str
    subject: str
    medium: str = "english"
    board: str = "BSE-TG"

    def _rows(self, content: bytes) -> list[dict[str, Any]]:
        book = load_workbook(
            io.BytesIO(content), read_only=True, data_only=True
        )
        wanted = {
            self.grade.strip().casefold(),
            f"grade {self.grade.strip()}".casefold(),
        }
        sheet = next(
            (n for n in book.sheetnames if n.strip().casefold() in wanted), None
        )
        if sheet is None:
            raise ValueError(
                f"No sheet for grade '{self.grade}'. Sheets in the workbook: "
                f"{', '.join(book.sheetnames)}."
            )
        rows = book[sheet].iter_rows(values_only=True)
        header = [
            "_".join(cell_text(h).lower().split()) for h in next(rows, ())
        ]
        if missing := [h for h in XLSX_HEADERS if h not in header]:
            raise ValueError(
                f"Sheet '{sheet}' has no column for {', '.join(missing)}. "
                f"Columns found: {', '.join(h for h in header if h)}."
            )
        return [dict(zip(header, row)) for row in rows]

    def _chapters(self, rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
        last: dict[str, Any] = {}
        periods: dict[tuple, dict[Any, dict[str, list[str]]]] = {}
        stray: list[int] = []
        for position, row in enumerate(rows, start=2):
            for key in ("chapter_number", "chapter_title", "period"):
                if cell_text(row[key]) == "":
                    row[key] = last.get(key)
                last[key] = row[key]
            if row["chapter_number"] is None:
                if cell_text(row["topic"]) or cell_text(
                    row["learning_outcomes"]
                ):
                    stray.append(position)
                continue
            group = periods.setdefault(
                (
                    whole_number(row["chapter_number"]),
                    cell_text(row["chapter_title"]),
                ),
                {},
            ).setdefault(
                whole_number(row["period"]),
                {"topics": [], "learning_outcomes": []},
            )
            for field, key in (
                ("topics", "topic"),
                ("learning_outcomes", "learning_outcomes"),
            ):
                if text := cell_text(row[key]):
                    group[field].append(text)
        if stray:
            raise ValueError(
                f"{len(stray)} rows have a topic or learning outcome but come "
                "before the first chapter number. Workbook rows, counting the "
                f"header as row 1: {stray}. Add the chapter number to these "
                "rows."
            )
        for (chapter_no, _), groups in periods.items():
            for period, group in groups.items():
                if group["learning_outcomes"] and not group["topics"]:
                    raise ValueError(
                        f"Chapter {chapter_no}, period {period} has learning "
                        "outcomes but no topic. Add the topic to the workbook."
                    )
        return [
            {
                "board": self.board,
                "medium": self.medium.lower(),
                "grade": re.sub(r"(?i)^grade\s*", "", self.grade).strip(),
                "subject": self.subject,
                "no": chapter_no,
                "name": name,
                "topic_groups": [
                    group
                    for _, group in sorted(
                        groups.items(), key=lambda kv: order(kv[0])
                    )
                    if group["topics"]
                ],
            }
            for (chapter_no, name), groups in sorted(
                periods.items(), key=lambda kv: order(kv[0][0])
            )
        ]

    async def run(self, ctx: IngestionContext) -> StepResult:
        if not self.grade or not self.subject:
            return StepResult(
                status=StepStatus.FAILURE,
                error=(
                    "An xlsx input needs the grade and subject parameters. "
                    "Add --param grade=6 --param subject=social."
                ),
            )
        try:
            chapters = self._chapters(self._rows(ctx.items[0].raw_content))
        except ValueError as e:
            return StepResult(status=StepStatus.FAILURE, error=str(e))
        ctx.metadata["period_chapters"] = chapters
        return StepResult(
            status=StepStatus.SUCCESS, metadata={"chapters": len(chapters)}
        )


register_step("xlsx_chapters", XlsxChapters, override=True)
register_step("section_graph", SectionGraph, override=True)
